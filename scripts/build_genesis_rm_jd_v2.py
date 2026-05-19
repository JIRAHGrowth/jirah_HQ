"""
Build the v2 Regional Manager - Metro Vancouver JD as a docx,
branded to match v1's formatting (Calibri, 18pt centered title,
12pt bold all-caps section headers, List Bullet style, same margins).

Output:
  [Genesis]/07 - Deliverables/Drafts/Regional-Manager-Metro-Vancouver-v2.docx
"""

from pathlib import Path
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- path guard (per CLAUDE.md scripts/_paths convention) ---
WS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WS / "scripts"))
try:
    from _paths import assert_no_legacy_segment  # type: ignore
except ImportError:
    def assert_no_legacy_segment(p: Path) -> None:
        if "NEW - JIRAH MASTER" in str(p):
            raise RuntimeError(f"forbidden legacy segment in path: {p}")

ONEDRIVE = Path(r"C:\Users\joshu\OneDrive - jirahgrowth.consulting")
OUT_DIR = (
    ONEDRIVE
    / "JIRAH Growth Partners - Shared"
    / "01 - Clients"
    / "Active"
    / "Genesis Systems"
    / "07 - Deliverables"
    / "Drafts"
)
OUT_PATH = OUT_DIR / "Regional-Manager-Metro-Vancouver-v2.docx"
assert_no_legacy_segment(OUT_PATH)


# ---------------- formatting helpers ----------------

EM = "—"  # em dash
EN = "–"  # en dash


def set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "BFBFBF", sz: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), sz)
        b.set(qn("w:color"), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(18)


def subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def section(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(12)


def subsection(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(11)


def body(doc: Document, text: str, justify: bool = True) -> None:
    """Body paragraph supporting **bold** segments."""
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_runs(p, text)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    _add_runs(p, text)


def _add_runs(paragraph, text: str) -> None:
    """Parse **bold** markers and render runs at Calibri 10.5pt."""
    parts = text.split("**")
    bold_flag = False
    for chunk in parts:
        if chunk == "":
            bold_flag = not bold_flag
            continue
        r = paragraph.add_run(chunk)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.bold = bold_flag
        bold_flag = not bold_flag


def hrule(doc: Document) -> None:
    """Thin horizontal rule via a bottom-border paragraph."""
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BFBFBF")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def make_table(doc: Document, headers: list[str], rows: list[list[str]],
               col_widths_in: list[float] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = False
    if col_widths_in:
        for i, w in enumerate(col_widths_in):
            for cell in table.columns[i].cells:
                cell.width = Inches(w)

    # header row
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, "1F2937")
        set_cell_borders(cell, color="1F2937", sz="6")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # data rows
    for ri, row in enumerate(rows):
        zebra = (ri % 2 == 1)
        for ci, val in enumerate(row):
            cell = table.rows[1 + ri].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            _add_runs(p, val)
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.name = "Calibri"
            if zebra:
                set_cell_shading(cell, "F4F1EA")
            set_cell_borders(cell, color="D9D2C4", sz="4")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


# ---------------- build doc ----------------

def build() -> Path:
    doc = Document()

    # Page setup matched to v1
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = Inches(1.1)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

    # Default font on Normal style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # Footer "Page X of Y"
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run("Page ")
    fr.font.size = Pt(9)
    # PAGE field
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    fr2 = footer_p.add_run()
    fr2._r.append(fld_begin); fr2._r.append(instr); fr2._r.append(fld_end)
    footer_p.add_run(" of ").font.size = Pt(9)
    fld_begin2 = OxmlElement("w:fldChar"); fld_begin2.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText"); instr2.text = "NUMPAGES"
    fld_end2 = OxmlElement("w:fldChar"); fld_end2.set(qn("w:fldCharType"), "end")
    fr3 = footer_p.add_run()
    fr3._r.append(fld_begin2); fr3._r.append(instr2); fr3._r.append(fld_end2)

    # === HEADLINE ===
    title(doc, f"REGIONAL MANAGER {EM} METRO VANCOUVER")
    subtitle(doc, "Genesis Controls  |  Senior Operational Leadership Track  |  v2")

    # === ABOUT THE ROLE ===
    section(doc, "About the Role")
    body(doc,
         "The Regional Manager " + EM + " Metro Vancouver carries single-point accountability for the people, "
         "projects, and performance of Genesis's Metro Vancouver office. The role reports directly to the Owner "
         "and is the operational bridge between Kelowna headquarters and the Metro Vancouver team " + EM + " "
         "translating company direction into regional execution, and bringing regional realities back into "
         "company-wide decisions.")
    body(doc,
         "This is the **first regional leadership role in Genesis's history**. The structure, cadence, and "
         "standards established here become the prototype for Alberta and any future regional offices. Strong "
         "performance positions the incumbent for broader operational, multi-region, or senior leadership "
         "responsibility as Genesis grows past the $10M revenue milestone.")

    # === ABOUT GENESIS CONTROLS ===
    section(doc, "About Genesis Controls")
    body(doc,
         "Founded in 2013 inside a large electrical contracting firm and operating as an independent entity since "
         "2017, Genesis Controls is one of British Columbia's fastest-growing controls and building integration "
         "companies. With expansion now active in Metro Vancouver and Washington State, Genesis serves hi-rise, "
         "commercial, institutional, residential, industrial, and healthcare clients. The firm's growth has been "
         "built on long relationships with general contractors, consulting engineers, and building owners across BC.")

    section(doc, "Core Values")
    body(doc,
         "Outstanding customer service through professional excellence, guaranteed workmanship, innovative "
         "technology, competitive pricing, and efficient scheduling " + EM + " held to the highest standards of "
         "safety and environmental protection. Integrity, dedication, and ethical practice with clients, "
         "suppliers, and staff.")

    # === REPORTING & RELATIONSHIPS ===
    section(doc, "Reporting & Relationships")
    body(doc, "**Reports to:** Owner / Chief Executive Officer.", justify=False)
    body(doc,
         "**Direct reports:** Project Managers, Controls Technologists, Controls Electricians (Journeyperson and "
         "Apprentice), and the Service Technician(s) assigned to Metro Vancouver.",
         justify=False)
    subsection(doc, "Headquarters peer functions (Kelowna-based, supporting all regions; do not report through this role)")
    bullet(doc, "**Office Manager** " + EM + " finance, A/P and A/R, payroll, group benefits, HR administration.")
    bullet(doc, "**Lead Estimator** " + EM + " bid strategy, subcontractor relationships, and estimating workload across regions.")
    bullet(doc, "**Technical PM** " + EM + " technical sequence-of-operations review, programming oversight, and cross-region engineering quality. The Technical PM is a **technical resource** to the Regional Manager, not an escalation path for management decisions. Management and resource concerns escalate directly to the Owner.")
    subsection(doc, "Cross-region peers (as established)")
    bullet(doc, "Future regional leaders (Alberta and beyond). Cross-region staff transfers, resource borrowing, and shared standards are jointly negotiated with counterpart regional managers.")

    # === OWNER CADENCE & SUPPORT ===
    section(doc, "Owner Cadence & Support Commitments")
    body(doc,
         "The Regional Manager " + EM + " Metro Vancouver is a new role in a growing firm. The Owner's "
         "commitments to the incumbent are part of the role's design, not optional add-ons.")
    bullet(doc, "**Weekly 1:1 with the Owner** " + EM + " 45 minutes, recurring, agenda-driven.")
    bullet(doc, "**Every-six-week Owner visits to Metro Vancouver** " + EM + " minimum two full days onsite per visit.")
    bullet(doc, "**Monthly regional business review** " + EM + " Regional Manager presents performance against KPI scorecard.")
    bullet(doc, "**Daily cross-office stand-up** " + EM + " 9:00" + EN + "9:15 internal regional team; 9:15" + EN + "9:30 cross-office video sync with Kelowna leadership.")
    bullet(doc, "**Quarterly leadership review** " + EM + " onsite at Kelowna headquarters with the full leadership team.")
    bullet(doc, "**Owner's standing first question:** " + chr(0x201C) + "How can I support you in this role? What do you need from me?" + chr(0x201D) + " The same question is the Regional Manager's standing first question to each direct report. This is a Genesis cultural primitive.")

    # === KEY RESPONSIBILITIES ===
    section(doc, "Key Responsibilities")

    subsection(doc, "People & Culture")
    bullet(doc, "Lead, coach, and develop the Metro Vancouver team across project management, technologist, trade, and service functions.")
    bullet(doc, "Own regional hiring, onboarding, and performance conversations in partnership with the Owner and Office Manager.")
    bullet(doc, "Build and sustain a strong regional culture consistent with Genesis's company-wide values, with particular attention to field-based and distributed team members.")
    bullet(doc, "Serve as the primary leadership presence for regional staff and the day-to-day bridge to headquarters for questions, concerns, and escalations.")
    bullet(doc, "Conduct quarterly 1:1 development conversations with each direct report; deliver a formal annual performance review by April 30 each year.")

    subsection(doc, "Operations & Delivery")
    bullet(doc, "Oversee execution of all projects and service work delivered from Metro Vancouver, ensuring quality, schedule, and budget commitments are met.")
    bullet(doc, "Enforce and continuously improve Genesis's documented SOPs " + EM + " bid-to-kickoff, project handoff, service dispatch, and project close-out " + EM + " at the regional level.")
    bullet(doc, "Coordinate with the Lead Estimator and Project Managers on project kickoff, scope clarity, and clean handoffs into field execution.")
    bullet(doc, "Own jobsite safety, provincial regulatory compliance, and a measurable safety culture across the region. **Specific accountability for the WCB claim trend** within the Metro Vancouver crew.")
    bullet(doc, "**Field-work boundary:** This role does not carry a dedicated service route, a routine PM workload, or scheduled on-tools field work. Field response is exception-only " + EM + " leadership presence at major escalations, customer-recovery situations, or safety incidents. Sustained hands-on field involvement (>10% of any week) is a leading indicator of operational structure failure and is to be raised in the next 1:1 with the Owner.")

    subsection(doc, "Client & Market")
    bullet(doc, "Maintain relationships with Metro Vancouver general contractors, consulting engineers, building owners, and service clients.")
    bullet(doc, "Represent Genesis at regional industry events (MCA, VRCA, and similar) and act as the senior face of the company in the region.")
    bullet(doc, "Carry **existing customer relationship development** " + EM + " deepening, expanding, and renewing existing accounts.")
    bullet(doc, "This role does **not** carry net-new business development or cold pipeline generation. Net-new BD remains with the Owner and the Business Development Manager.")
    bullet(doc, "Identify regional market trends, competitive activity, and growth opportunities; surface them monthly to the Owner.")

    subsection(doc, "Financial & Reporting")
    bullet(doc, "Own regional operational performance against agreed KPIs (see Success Metrics).")
    bullet(doc, "Monitor regional job costing in partnership with the Office Manager and Project Managers; flag margin erosion, scope creep, and cash flow risks early.")
    bullet(doc, "Prepare the **monthly regional performance summary** for the Owner and contribute to the quarterly leadership review.")
    bullet(doc, "Provide input to annual regional budgeting and workforce planning.")

    # === AUTHORITY MATRIX ===
    section(doc, "Authority Matrix")
    body(doc,
         "The Regional Manager carries the following authority. **Numerical thresholds will be confirmed in the "
         "May 5 sit-down and entered here within 14 days. The structure is locked; the numbers are the live variable.**")

    auth_rows = [
        ["Day-to-day scheduling, project assignment, field resource allocation",
         "Full",
         "None within region"],
        ["Operational purchases (parts, tools, materials, subcontractor services)",
         f"Up to **${'{X}'}** per single transaction",
         "> threshold " + EM + " Owner"],
        ["Hiring within approved headcount " + EM + " trade, technologist, service",
         "Full",
         "New roles outside plan " + EM + " Owner"],
        ["Termination & disciplinary decisions, regional staff",
         "Joint with Owner",
         "All cases consult Owner first"],
        ["Project commitments / contract authority",
         f"Up to **${'{X}'}** project value",
         "> threshold " + EM + " Owner"],
        ["Subcontractor engagement on regional projects",
         "Full, within Lead Estimator's pricing/scope",
         "Outside framework " + EM + " joint with Lead Estimator"],
        ["Client issue resolution, service recovery, customer concessions",
         f"Up to **${'{X}'}** per incident",
         "> threshold " + EM + " Owner"],
        ["Cross-region staff borrowing",
         "Joint with counterpart regional leader",
         "None"],
        ["Bid strategy, go/no-go",
         "Joint with Lead Estimator",
         "None"],
        ["Headcount changes beyond plan, executive hires, senior compensation",
         "Owner-owned",
         "All " + EM + " Regional Manager recommends"],
        ["Pricing changes, service packaging, commercial policy",
         "Owner-owned",
         "All"],
        ["Vehicle purchases, office leases, major capital",
         "Owner-owned",
         "All"],
        ["Material legal, financial, or reputational exposure",
         "Owner-owned",
         "Immediate notification"],
    ]
    make_table(
        doc,
        ["Decision Area", "Regional Manager Authority", "Escalation Trigger"],
        auth_rows,
        col_widths_in=[3.0, 2.0, 1.7],
    )

    # === SUCCESS METRICS ===
    section(doc, "Success Metrics")
    body(doc, "Regional performance is reviewed monthly against:")
    bullet(doc, "Regional gross margin against target.")
    bullet(doc, "Project on-time and on-budget completion rate.")
    bullet(doc, "Service response time (request to dispatch) and first-visit resolution rate.")
    bullet(doc, "Regional employee retention and voluntary turnover.")
    bullet(doc, "Field utilization and overtime trend.")
    bullet(doc, "Safety incident rate, WCB claim trend, and near-miss reporting participation.")
    bullet(doc, "Regional revenue contribution toward Genesis's $10M company-wide milestone.")

    # === TRANSITION ARC ===
    section(doc, "Transition Arc " + EM + " First 180 Days")
    body(doc,
         "This role is being filled by an internal promotion. The first 180 days are a structured transition with "
         "explicit milestones, not steady-state operations.")

    subsection(doc, "First 30 days")
    bullet(doc, "Daily Vancouver stand-up running without skipping.")
    bullet(doc, "Cross-office video stand-up running at least twice weekly.")
    bullet(doc, "Whiteboard installed; first regional planning grid drawn (3-month / 6-month).")
    bullet(doc, "1:1 conducted with each direct report using the " + chr(0x201C) + "How can I support you?" + chr(0x201D) + " frame.")
    bullet(doc, "Service-site successor running two service sites independently.")
    bullet(doc, "Field hands-on time reduced from baseline to under 20% of week.")

    subsection(doc, "Days 30" + EN + "60")
    bullet(doc, "Weekly stand-up cadence locked.")
    bullet(doc, "First independent monthly regional business review delivered to the Owner.")
    bullet(doc, "Underperformer decisions made (transition out or formalized expectations).")
    bullet(doc, "Major service contracts day-to-day work fully transitioned to service successor.")
    bullet(doc, "Service-site institutional knowledge transferred to successor in writing " + EM + " per-site equipment lists, sequence-of-operation history, prior issues.")

    subsection(doc, "Days 60" + EN + "90")
    bullet(doc, "Progress billing cycle owned end-to-end for one full Vancouver cycle.")
    bullet(doc, "Material ordering and supplier relationships transferred from Owner to Regional Manager.")
    bullet(doc, "Q3 regional hiring plan presented to Owner as a **recommendation**, not a question.")
    bullet(doc, "Compensation structure finalized and signed.")

    subsection(doc, "Days 90" + EN + "180")
    bullet(doc, "First quarterly Regional Manager scorecard delivered.")
    bullet(doc, "Existing-customer relationship plan executed (industry conference attended jointly with Owner).")
    bullet(doc, "Field hands-on time at under 10% of week; firefighting reframed as exception-only.")
    bullet(doc, "Six-month role confirmation review (see Performance Review section).")

    subsection(doc, "What this role stops doing " + EM + " explicitly")
    bullet(doc, "Sole-source crew assignment (becomes a Friday stand-up artifact).")
    bullet(doc, "Day-to-day routine service work on the major Vancouver service contracts.")
    bullet(doc, "Routine field service rotation.")
    bullet(doc, "Acting as the default escalation for any single major customer relationship.")

    # === COMPENSATION FRAMEWORK ===
    section(doc, "Compensation Framework")
    bullet(doc, "**Structure:** Base salary plus performance variable.")
    bullet(doc, "**Variable tied to:** regional gross margin against target, regional employee retention, safety performance, project on-time/on-budget rate. Specific weighting confirmed at structure finalization.")
    bullet(doc, "**Vehicle / phone / professional development:** included; specifics confirmed at structure finalization.")
    bullet(doc, "**Review cadence:** annual review in April; informal mid-year touch in October.")
    bullet(doc, "**Long-term incentive eligibility:** as Genesis formalizes a multi-year incentive plan, this role is eligible.")
    bullet(doc, "**Final compensation structure locked:** within 14 days of role acceptance. The structure is in this document; the numbers are confirmed in a dedicated compensation conversation.")

    # === PERFORMANCE REVIEW & ROLE CONFIRMATION ===
    section(doc, "Performance Review & Role Confirmation")
    bullet(doc, "**30-day check-in** " + EM + " Owner and Regional Manager review transition deliverables.")
    bullet(doc, "**90-day check-in** " + EM + " Owner and Regional Manager review the full transition arc against milestones.")
    bullet(doc, "**Six-month role confirmation review** " + EM + " formal review at the end of October. Outcome is one of: role confirmed (default expectation); role confirmed with development plan; role not yet confirmed " + EM + " extended review. This is a fairness mechanism, not a probationary trap. Both parties walk in with prepared input.")
    bullet(doc, "**Annual performance review** " + EM + " every April, against the full Success Metrics scorecard.")

    # === QUALIFICATIONS ===
    section(doc, "Qualifications & Experience")

    subsection(doc, "Required")
    bullet(doc, "Minimum eight years of operational experience in building controls, BMS, mechanical contracting, or an adjacent construction trade, with at least three years in a formal supervisory or project management role.")
    bullet(doc, "Demonstrated ability to lead cross-disciplinary teams across project management, technologist, trade, and service functions.")
    bullet(doc, "Strong working knowledge of HVAC systems, building automation, and BC electrical and building codes and safety regulations.")
    bullet(doc, "Ability to read and interpret construction documents, project schedules, and financial statements at a P&L line-item level.")
    bullet(doc, "Excellent written and verbal communication, with the ability to translate between field, engineering, and executive audiences.")
    bullet(doc, "Valid BC driver's license and access to reliable transportation.")

    subsection(doc, "Preferred")
    bullet(doc, "Journeyperson Electrician certification or equivalent controls-trade credential.")
    bullet(doc, "Post-secondary education in construction management, engineering, business, or a related field.")
    bullet(doc, "Prior experience managing a geographically distinct office or region.")
    bullet(doc, "Familiarity with project management and job-costing platforms (SiteMax, Procore, Sage, or equivalent).")
    bullet(doc, "Track record of culture-building in a field-based, distributed team.")

    # === BEHAVIORAL COMPETENCIES ===
    section(doc, "Behavioral Competencies")
    body(doc, "These behaviors are expected in role and form the spine of every coaching 1:1.")
    comp_rows = [
        ["**Coaches in the moment**",
         "Walks a tech through a sequence-of-operations question instead of solving it himself. Names the teaching moment out loud.",
         chr(0x201C) + "It's faster if I just do it." + chr(0x201D) + " Default to fixing problems personally."],
        ["**Builds operating cadence**",
         "Standing meetings actually run on the calendar. Decisions tracked. Tomorrow's plan visible Friday.",
         "Reactive scheduling. Fires drive the day. Stand-ups skipped when " + chr(0x201C) + "things are busy." + chr(0x201D)],
        ["**Names the hard conversation early**",
         "Performance issues surfaced in the first 1:1 with Owner, with a recommended path. WCB claim risk raised proactively.",
         "Sits on tough calls. Lets a perception or conflict run for weeks before naming it."],
        ["**Escalates up, not sideways**",
         "Resource conflict goes to the Owner directly. Technical question goes to Technical PM. The two paths are distinct and named.",
         "Escalates management problems sideways to a peer in another office."],
        ["**Owner-mindset**",
         "Asks " + chr(0x201C) + "what's the regional contribution to the $10M target this quarter?" + chr(0x201D) + " before being asked. Brings recommendations, not questions.",
         "Walks into Owner meetings with open issues to be solved by someone else."],
        ["**Protects team capacity**",
         "Says no to overflow when the math doesn't work. Pushes back on commitment dates that would force overtime.",
         "Absorbs every additional ask. " + chr(0x201C) + "We'll figure it out." + chr(0x201D)],
        ["**Calm under operational pressure**",
         "Service emergency lands; the team sees the same demeanor as on a normal Tuesday.",
         "Visible reactivity. Pulls senior people into firefighting by tone."],
    ]
    make_table(
        doc,
        ["Competency", "Looks Like", "Does Not Look Like"],
        comp_rows,
        col_widths_in=[1.7, 2.5, 2.5],
    )

    # === WORK ENVIRONMENT & TRAVEL ===
    section(doc, "Work Environment & Travel")
    bullet(doc, "Based at the Metro Vancouver office.")
    bullet(doc, "Regular jobsite presence across Metro Vancouver and the Fraser Valley " + EM + " exception-led, not routine.")
    bullet(doc, "**Monthly travel to Kelowna headquarters** for leadership reviews and team rhythm " + EM + " typically two to three days per visit.")
    bullet(doc, "Quarterly extended leadership offsite (one week per quarter at Kelowna).")
    bullet(doc, "Occasional travel to other Genesis regional offices as the company expands.")

    # === GROWTH PATHWAY ===
    section(doc, "Growth Pathway")
    body(doc,
         "The Regional Manager " + EM + " Metro Vancouver role sits on Genesis's senior operational leadership "
         "track and is the prototype for the Alberta region and beyond. Strong performance positions the incumbent "
         "for broader operational scope (multi-region oversight), senior leadership responsibility, or expanded "
         "ownership of company-wide operating systems as Genesis scales.")
    body(doc,
         "The role's named successor for the regional service function is the Service Technician identified during "
         "the transition arc.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    p = build()
    print(f"Wrote: {p}")
