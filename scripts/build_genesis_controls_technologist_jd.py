"""
Build the v1 Controls Technologist JD as a docx, matching the locked
Genesis template (Calibri, 18pt centered title, 12pt bold all-caps section
headers, List Bullet style, navy header tables with zebra striping).

Source material:
  - Controls-Technologist.pdf (Genesis current JD on file) — primary source
  - Genesis template (RM-MV-Final, Lead-Estimator-v1, Senior-Controls-Technologist-v1)

Tier: mid IC. Incumbents move into this seat from the Junior tier; the
seat reports into the Senior Controls Technologist (technical line) and
the regional Project Manager (delivery line). No direct reports. Genesis
uses Engineer / Technologist / Programmer interchangeably; "Controls
Technologist" is the canonical title for this seat.

Output:
  [Genesis]/07 - Deliverables/Drafts/Controls-Technologist-v1.docx
"""

from pathlib import Path
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WS / "scripts"))
from _paths import assert_no_legacy_segment, ACTIVE_CLIENTS  # type: ignore

OUT_DIR = ACTIVE_CLIENTS / "Genesis Systems" / "07 - Deliverables" / "Drafts"
OUT_PATH = OUT_DIR / "Controls-Technologist-v1.docx"
assert_no_legacy_segment(OUT_PATH)

EM = "—"
EN = "–"
LDQ = chr(0x201C)
RDQ = chr(0x201D)


def set_cell_shading(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="BFBFBF", sz="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), sz)
        b.set(qn("w:color"), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(18)


def subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(12)


def subsection(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(11)


def body(doc, text, justify=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_runs(p, text)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    _add_runs(p, text)


def _add_runs(paragraph, text):
    parts = text.split("**")
    bold_flag = False
    for chunk in parts:
        if chunk == "":
            bold_flag = not bold_flag
            continue
        r = paragraph.add_run(chunk)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.bold = bold_flag
        bold_flag = not bold_flag


def make_table(doc, headers, rows, col_widths_in=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = False
    if col_widths_in:
        for i, w in enumerate(col_widths_in):
            for cell in table.columns[i].cells:
                cell.width = Inches(w)

    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, "1F2937")
        set_cell_borders(cell, color="1F2937", sz="6")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for ri, row in enumerate(rows):
        zebra = (ri % 2 == 1)
        for ci, val in enumerate(row):
            cell = table.rows[1 + ri].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            _add_runs(p, val)
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.name = "Calibri"
            if zebra:
                set_cell_shading(cell, "F4F1EA")
            set_cell_borders(cell, color="D9D2C4", sz="4")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def build():
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = Inches(1.1)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # Footer: Page X of Y
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run("Page ")
    fr.font.size = Pt(9)
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    fr2 = footer_p.add_run()
    fr2._r.append(fld_begin); fr2._r.append(instr); fr2._r.append(fld_end)
    footer_p.add_run(" of ").font.size = Pt(9)
    fld_begin2 = OxmlElement("w:fldChar"); fld_begin2.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText"); instr2.text = "NUMPAGES"
    fld_end2 = OxmlElement("w:fldChar"); fld_end2.set(qn("w:fldCharType"), "end")
    fr3 = footer_p.add_run()
    fr3._r.append(fld_begin2); fr3._r.append(instr2); fr3._r.append(fld_end2)

    # === HEADLINE ===
    title(doc, "CONTROLS TECHNOLOGIST")
    subtitle(doc, "Genesis Building Controls  |  Engineering & Programming Function")

    # === ABOUT THE ROLE ===
    section(doc, "About the Role")
    body(doc,
         "The Controls Technologist is a talented and experienced individual responsible for designing, "
         "implementing, and maintaining HVAC control systems that ensure optimal performance, energy "
         "efficiency, and comfort for Genesis's clients. The role works closely with project managers, design "
         "engineers, technicians, and clients to deliver high-quality solutions that meet project requirements "
         "and exceed customer expectations.")
    body(doc,
         "The Controls Technologist sits in the **middle of Genesis's controls-technologist track**. The seat "
         "reports into the Senior Controls Technologist on the technical line and the regional Project Manager "
         "on the delivery line for day-to-day project work. Junior Controls Technologists sit below; the Senior "
         "Controls Technologist sits above. Career progression is into the senior seat.")
    body(doc,
         "Genesis treats Engineer, Technologist, and Programmer as functionally synonymous titles within the "
         "controls trade. The work is the same; the title reflects the seat's position in the bench rather than "
         "a separate function.")

    # === ABOUT GENESIS BUILDING CONTROLS ===
    section(doc, "About Genesis Building Controls")
    body(doc,
         "Founded in 2013 inside a large electrical contracting firm and operating as an independent entity "
         "since 2017, Genesis Building Controls is one of British Columbia's fastest-growing controls and "
         "building integration companies. With expansion now active in Metro Vancouver and preparing for "
         "Alberta, Genesis serves hi-rise, commercial, institutional, residential, industrial, and healthcare "
         "clients. The firm's growth has been built on long relationships with general contractors, consulting "
         "engineers, and building owners across BC.")

    section(doc, "Core Values")
    body(doc,
         "Outstanding customer service through professional excellence, guaranteed workmanship, innovative "
         "technology, competitive pricing, and efficient scheduling " + EM + " held to the highest standards "
         "of safety and environmental protection. Integrity, dedication, and ethical practice with clients, "
         "suppliers, and staff.")

    # === REPORTING & RELATIONSHIPS ===
    section(doc, "Reporting & Relationships")
    body(doc,
         "**Reports to:** Senior Controls Technologist (technical line " + EM + " standards, code review, "
         "technical escalation, professional development) and the regional Project Manager (delivery line " + EM +
         " project-specific scheduling, client coordination, scope clarity).",
         justify=False)
    body(doc,
         "**Direct reports:** None. The Controls Technologist provides on-the-job guidance to Junior Controls "
         "Technologists and apprentices on assigned projects without carrying formal supervisory accountability.",
         justify=False)

    subsection(doc, "Primary internal partners")
    bullet(doc, "**Senior Controls Technologist** " + EM + " technical standards, sequence-of-operations review, code review on submitted work, technical escalation point, quarterly development conversations.")
    bullet(doc, "**Project Manager(s)** " + EM + " day-to-day project delivery, scheduling, client coordination, project-side escalation.")
    bullet(doc, "**Junior Controls Technologists** " + EM + " on-the-job guidance, knowledge transfer, paired commissioning support on assigned projects.")
    bullet(doc, "**Lead Estimator / Estimators** " + EM + " bid-stage scope input on assigned pursuits, site-survey and integration estimates.")
    bullet(doc, "**Controls Electricians (Forepersons, Journeypersons, Apprentices)** " + EM + " installation coordination, field commissioning support.")
    bullet(doc, "**Service & Warranty Coordinator** " + EM + " service-side programming, retrofit programming, post-handover defect resolution on assigned projects.")
    bullet(doc, "**Office Manager** " + EM + " software license, PD support, contract administration touchpoints.")

    # === CADENCE & SUPPORT ===
    section(doc, "Cadence & Support Commitments")
    body(doc,
         "The Controls Technologist is supported by a defined cadence with both reporting lines. Every "
         "manager and the bench lead operate on a recurring rhythm so blockers, technical questions, and "
         "development conversations have a defined home, not an ad-hoc scramble.")
    bullet(doc, "**Weekly 1:1 with the Senior Controls Technologist** " + EM + " 30 minutes, recurring, agenda-driven. Active project work, blockers, code-review feedback, professional development.")
    bullet(doc, "**Weekly engineering sync** " + EM + " full controls-technologist pool with the Senior Controls Technologist and Technical Lead " + EM + " standards updates, knowledge transfer, sequence-of-operations review.")
    bullet(doc, "**Project review with assigned PM(s)** " + EM + " cadence calibrated to project rhythm; standing project-side check-in for delivery, scheduling, and client coordination.")
    bullet(doc, "**Quarterly development conversation** with the Senior Controls Technologist " + EM + " progression on the technical ladder, certification path, project-mix development.")
    bullet(doc, "**Standing first question:** " + LDQ + "How can I support you in this role? What do you need from me?" + RDQ + " The Senior Controls Technologist asks it of the Controls Technologist; the Controls Technologist brings the answer to the next 1:1. This is a Genesis cultural primitive.")

    # === KEY RESPONSIBILITIES ===
    section(doc, "Key Responsibilities")

    subsection(doc, "HVAC Controls Design & Engineering")
    bullet(doc, "Design and develop HVAC control systems based on project specifications, building codes, and industry standards.")
    bullet(doc, "Collaborate with project managers, design engineers, and clients to understand project requirements and objectives.")
    bullet(doc, "Perform detailed site surveys and assessments to gather the information necessary for system design and implementation.")
    bullet(doc, "Select appropriate control components, sensors, actuators, and other hardware necessary for system integration.")
    bullet(doc, "Develop control strategies and sequences of operation to optimize HVAC system performance and energy efficiency.")
    bullet(doc, "Create control system drawings, schematics, and diagrams using CAD software.")

    subsection(doc, "Programming & Implementation")
    bullet(doc, "Program and configure building automation systems (BAS) and control devices such as PLCs, DDC controllers, and thermostats.")
    bullet(doc, "Work to Genesis's controls programming standards " + EM + " sequence-of-operations templates, controller programming conventions, BACnet object naming, graphics, alarming, and trending discipline. Deviations are flagged to the Senior Controls Technologist, not improvised.")
    bullet(doc, "Submit programming work through code review with the Senior Controls Technologist on schedule and to standard. Code review is the developmental spine of the seat " + EM + " brought into it as collaboration, not gatekeeping.")
    bullet(doc, "Maintain reusable logic blocks and design patterns developed on assigned projects in Genesis's controls library, so the next technologist to face a similar project starts ahead.")

    subsection(doc, "Commissioning & Field Support")
    bullet(doc, "Conduct system testing, commissioning, and troubleshooting to ensure proper functionality and performance.")
    bullet(doc, "Provide technical support and assistance to field technicians and subcontractors during installation, startup, and maintenance activities.")
    bullet(doc, "Support post-commissioning defect resolution on assigned projects; surface root-cause patterns to the Senior Controls Technologist so standards can be updated where appropriate.")
    bullet(doc, "Walk assigned projects before programming starts on anything beyond routine scope " + EM + " catch mechanical realities the drawings don't show.")

    subsection(doc, "Coordination & Communication")
    bullet(doc, "Consult and communicate with engineers, architects, owners, contractors, and subcontractors at the technical level appropriate to the project.")
    bullet(doc, "Coordinate project tasks in partnership with the assigned Project Manager(s) " + EM + " scope clarity, scheduling, site-coordination, change management.")
    bullet(doc, "Stay current on advancements in HVAC controls technology, industry trends, and Genesis-approved platforms. Bring observations into the weekly engineering sync.")
    bullet(doc, "Provide on-the-job guidance to Junior Controls Technologists and apprentices on assigned projects " + EM + " paired commissioning, walk-throughs of sequence logic, code-review readiness.")

    # === AUTHORITY MATRIX ===
    section(doc, "Authority Matrix")
    body(doc,
         "Authority is calibrated to the seat's mid-IC tier. The Controls Technologist operates within "
         "Genesis's standards and approved platforms; outside-of-pattern decisions route to the Senior "
         "Controls Technologist.")

    auth_rows = [
        ["Day-to-day controls design and programming decisions within approved project scope",
         "Full",
         "Outside approved standards " + EN + " Senior Controls Technologist"],
        ["Component / hardware selection within approved vendor lists",
         "Full",
         "New vendor / non-standard component " + EN + " Senior Controls Technologist"],
        ["Sequence-of-operations interpretation when consulting-engineer drawings are ambiguous",
         "Advisory " + EM + " surface to Senior Controls Technologist",
         "All ambiguities " + EN + " Senior Controls Technologist before programming"],
        ["Platform / version selection and upgrade decisions",
         "None (works within approved platforms)",
         "All " + EN + " Senior Controls Technologist"],
        ["Project schedule and client commitments",
         "Coordinates with assigned PM",
         "Schedule slip or scope change " + EN + " PM + Senior Controls Technologist"],
        ["Field-escalation resolution within controls scope, approved methods",
         "Full",
         "Outside controls scope " + EN + " PM; novel integration " + EN + " Senior Controls Technologist"],
        ["Software license / capital tooling requests",
         "Advisory " + EM + " recommends to Senior Controls Technologist",
         "All " + EN + " Senior Controls Technologist + Office Manager"],
        ["External integration partner / contractor engagement",
         "None (routes through Senior Controls Technologist)",
         "All " + EN + " Senior Controls Technologist"],
        ["Material legal, financial, or reputational exposure",
         "Owner-owned",
         "Immediate notification " + EN + " Senior Controls Technologist + Regional Manager"],
    ]
    make_table(
        doc,
        ["Decision Area", "Controls Technologist Authority", "Escalation Trigger"],
        auth_rows,
        col_widths_in=[2.8, 2.2, 1.7],
    )

    # === SUCCESS METRICS ===
    section(doc, "Success Metrics")
    body(doc, "The Controls Technologist's performance is reviewed against:")
    bullet(doc, "**Project on-time / on-spec delivery** on assigned projects " + EM + " controls-side schedule and scope performance against PM-held targets.")
    bullet(doc, "**Commissioning defect rate** on assigned projects " + EM + " rework attributable to controls work. Trend over rolling quarter.")
    bullet(doc, "**Standards adoption** " + EM + " share of work shipped through Genesis's templates, library, and code-review process versus one-off improvisation.")
    bullet(doc, "**Code-review first-pass quality** " + EM + " share of submitted work that clears code review without substantive rework. Tracked as a development metric, not a punitive one.")
    bullet(doc, "**Field support responsiveness** " + EM + " turnaround on field escalations from assigned projects.")
    bullet(doc, "**Professional development progression** " + EM + " certification milestones, technical-ladder progression, project-mix breadth.")
    bullet(doc, "**Project Manager and Senior Controls Technologist feedback** at quarterly 1:1s.")

    # === COMPENSATION FRAMEWORK ===
    section(doc, "Compensation Framework")
    bullet(doc, "**Structure:** base salary plus performance variable.")
    bullet(doc, "**Variable tied to:** project on-time / on-spec performance, commissioning defect rate on assigned projects, and professional development progression. Specific weighting confirmed at structure finalization.")
    bullet(doc, "**Vehicle / phone / professional development:** vehicle (or vehicle allowance), phone, and PD budget included. PD allowance covers certifications (Tridium / Niagara, BACnet, integration platforms) and conference attendance on the Senior Controls Technologist's recommended development path.")
    bullet(doc, "**Review cadence:** annual, April; informal mid-year touch in October.")
    bullet(doc, "**Long-term incentive eligibility:** as Genesis formalizes a multi-year incentive plan, this role is eligible.")
    bullet(doc, "**Final compensation structure locked:** within 14 days of role acceptance.")

    # === PERFORMANCE REVIEW ===
    section(doc, "Performance Review & Role Confirmation")
    bullet(doc, "**30-day check-in** " + EM + " Senior Controls Technologist and Controls Technologist review onboarding deliverables, first project assignment, standards orientation.")
    bullet(doc, "**90-day check-in** " + EM + " Senior Controls Technologist and Controls Technologist review the first-quarter performance, first independent commissioning support, and code-review trajectory.")
    bullet(doc, "**Six-month role confirmation review** " + EM + " formal review at the end of the second quarter in seat. Outcome is one of: role confirmed (default expectation), role confirmed with development plan, role not yet confirmed " + EM + " extended review. This is a fairness mechanism, not a probationary trap. Both parties walk in with prepared input.")
    bullet(doc, "**Annual performance review** " + EM + " every April, against the full Success Metrics scorecard.")

    # === QUALIFICATIONS ===
    section(doc, "Qualifications & Experience")

    subsection(doc, "Required")
    bullet(doc, "Minimum 3 years of HVAC building controls experience or related field " + EM + " design, programming, commissioning.")
    bullet(doc, "Bachelor's degree in Mechanical Engineering, Electrical Engineering, or related field " + EM + " or equivalent technologist credential (CET / RET) with a demonstrable track record.")
    bullet(doc, "Strong working knowledge of HVAC systems " + EM + " air handling units, chillers, boilers, and VAV systems.")
    bullet(doc, "Experience with building automation systems (BAS) and control protocols such as BACnet, Modbus, and LonWorks.")
    bullet(doc, "Proficient in at least one programming language commonly used in HVAC controls " + EM + " ladder logic, C/C++, Java, Python, or equivalent.")
    bullet(doc, "Proficiency in CAD software for control system design and documentation.")
    bullet(doc, "Excellent problem-solving skills and attention to detail.")
    bullet(doc, "Effective communication skills, both verbal and written.")
    bullet(doc, "Ability to work independently and as part of a team in a fast-paced environment.")
    bullet(doc, "Valid BC driver's license and access to reliable transportation.")

    subsection(doc, "Preferred")
    bullet(doc, "Tridium Niagara certifications (TCP or equivalent).")
    bullet(doc, "Master's degree in Mechanical or Electrical Engineering.")
    bullet(doc, "Professional certifications such as Certified Energy Manager (CEM) or Certified Automation Professional (CAP).")
    bullet(doc, "Experience with the specific controller stacks Genesis works in (Distech, Honeywell, Johnson Controls, Schneider Electric).")
    bullet(doc, "Familiarity with the BC consulting engineer and general contractor market.")
    bullet(doc, "Prior experience pairing with junior technologists or apprentices on commissioning days.")

    # === BEHAVIOURAL COMPETENCIES ===
    section(doc, "Behavioural Competencies")
    body(doc, "These are the behaviours expected in role and the spine of every coaching 1:1.")
    comp_rows = [
        ["**Writes for the next person who reads the code**",
         "Logic is readable. Names are honest. Graphics, alarming, and trending follow the standard. The next technologist to touch the project can find their footing without an interpreter.",
         "Clever shortcuts that compile and ship. " + LDQ + "I'll explain it later" + RDQ + " code that becomes load-bearing on the project."],
        ["**Works to the standard**",
         "Sequence-of-operations templates, BACnet object naming, graphics, alarming, and trending follow the firm standard. Deviations get surfaced and explained, not improvised.",
         "Each project gets a personal style. Standards treated as suggestions. Reusable logic re-derived because the library wasn't checked first."],
        ["**Brings the question early**",
         "Ambiguous drawing, missing assumption, or unfamiliar integration surfaced at the next 1:1 or sync " + EM + " not at commissioning. " + LDQ + "I'm not sure, let me check" + RDQ + " is a strength, not a weakness.",
         "Programs to the drawings as drawn and lets the gaps land at commissioning. Sits on uncertainty until it becomes a field problem."],
        ["**Calm under commissioning pressure**",
         "Commissioning day lands; the team sees the same demeanor as on a normal Tuesday. Solves the actual problem instead of the visible one.",
         "Visible reactivity under pressure. Pulls senior staff into firefighting by tone."],
        ["**Holds the line on sequence-of-operations integrity**",
         "Surfaces ambiguities before programming starts. Names what a building is being asked to do and what the drawings don't support.",
         LDQ + "We'll figure it out on site" + RDQ + " attitude. Treats sequence-of-ops review as paperwork rather than risk-reduction."],
        ["**Reads the field, not just the spec**",
         "Walks assigned jobs before programming. Knows what the controls have to do in a real building under real load. Notices what the drawings don't show.",
         "Treats the spec as gospel. Programs in isolation from the mechanical reality. Surprises at commissioning that an hour on site would have prevented."],
        ["**Coachable**",
         "Invites code review. Names what they don't know. Takes feedback as development input, not as critique. Asks for the project that stretches the bench.",
         "Defensive about code review. Treats feedback as performance management. Avoids the project that would expose a gap."],
    ]
    make_table(
        doc,
        ["Competency", "Looks Like", "Does Not Look Like"],
        comp_rows,
        col_widths_in=[1.7, 2.5, 2.5],
    )

    # === WORK ENVIRONMENT & TRAVEL ===
    section(doc, "Work Environment & Travel")
    bullet(doc, "Based at the assigned regional office (Kelowna headquarters or Metro Vancouver).")
    bullet(doc, "Regular jobsite presence across the assigned region for site surveys, programming, commissioning, and troubleshooting.")
    bullet(doc, "Occasional travel to the counterpart regional office for cross-regional commissioning support, training, or paired work with the Senior Controls Technologist.")
    bullet(doc, "Industry training and certification events " + EM + " Tridium / Niagara, BACnet, and equivalent " + EM + " on the development path agreed with the Senior Controls Technologist.")

    # === GROWTH PATHWAY ===
    section(doc, "Growth Pathway")
    body(doc,
         "The Controls Technologist seat sits on Genesis's technical track. The clear next step is the "
         "**Senior Controls Technologist** seat " + EM + " Genesis's apex technical-depth role across all "
         "regions, with cross-regional standards authority and bench leadership responsibility. Progression "
         "is paced by demonstrated breadth across controller families, project complexity, and standards-"
         "setting contribution to the controls library.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    p = build()
    print(f"Wrote: {p}")
