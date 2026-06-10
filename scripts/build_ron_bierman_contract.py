"""Build Ron Bierman's Senior Technical Advisor (Contractor) Agreement.

Mirrors the structure of Jason Lotoski's signed Senior Consultant Agreement
(Oct 2025) so JIRAH's legal voice and signing pattern stay consistent.
Adapts the IP, security, and handover clauses for a technical-build role
touching client production systems.

Output: 05 - Corporate / Legal & Agreements /
        Consultancy Services Agreement - Ron Bierman - June 2026.docx
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _paths import CORPORATE, assert_no_legacy_segment


BODY_FONT = "Calibri"
BODY_SIZE = Pt(11)
HEADING_SIZE = Pt(13)
TITLE_SIZE = Pt(18)


def set_run(run, *, bold=False, size=BODY_SIZE, font=BODY_FONT):
    run.font.name = font
    run.font.size = size
    run.bold = bold


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), bold=True, size=TITLE_SIZE)


def add_heading(doc, text):
    p = doc.add_paragraph()
    set_run(p.add_run(text), bold=True, size=HEADING_SIZE)


def add_subheading(doc, text):
    p = doc.add_paragraph()
    set_run(p.add_run(text), bold=True)


def add_para(doc, text, *, bold=False):
    p = doc.add_paragraph()
    set_run(p.add_run(text), bold=bold)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run(p.add_run(item))


def add_signature_block(doc, name, role_line):
    doc.add_paragraph()
    add_para(doc, "______________________________")
    add_para(doc, name, bold=True)
    add_para(doc, role_line)


def build(out_path: Path) -> None:
    assert_no_legacy_segment(out_path)

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE

    # ---------- Title & Parties ----------
    add_title(doc, "Consultancy Services Agreement")
    doc.add_paragraph()

    add_para(
        doc,
        'This Consultancy Services Agreement (the "Agreement") is made and '
        'entered into between:',
    )
    add_para(
        doc,
        'JIRAH Growth Consulting Ltd. (the "Company"), with its principal '
        "place of business in Kelowna, British Columbia; and",
    )
    add_para(
        doc,
        'Ron Bierman (the "Consultant"), an independent contractor, '
        "with their place of residence at ______________________________.",
    )
    add_para(
        doc,
        "Effective Date: June 10, 2026.",
        bold=True,
    )

    # ---------- 1. Term ----------
    add_heading(doc, "1. Term")
    add_para(
        doc,
        "This Agreement shall commence on the Effective Date and remain in "
        'force for an initial period of six (6) months (the "Initial Term"). '
        "Following the Initial Term, the Agreement shall automatically renew "
        "on a month-to-month basis unless either party provides thirty (30) "
        "days' written notice of non-renewal or termination. The Company may "
        "terminate this Agreement immediately for cause, including breach of "
        "this Agreement, breach of confidentiality or intellectual property "
        "obligations, security breach, misconduct, or conflict of interest.",
    )

    # ---------- 2. Scope of Work ----------
    add_heading(doc, "2. Scope of Work")
    add_para(
        doc,
        'The Consultant agrees to perform the services described in '
        'Appendix A (the "Services"). The Consultant\'s primary '
        "responsibility is providing senior technical leadership across "
        "JIRAH's AI product and operating-system builds — including "
        "architectural oversight, security and production-readiness review, "
        "infrastructure setup, and the handover of production-grade systems "
        "to JIRAH clients. The Consultant works in close collaboration with "
        "Jason Lotoski, who leads AI build cycles, and with Joshua Marshall "
        "on commercial and client-facing matters.",
    )

    # ---------- 3. Compensation ----------
    add_heading(doc, "3. Compensation")
    add_para(
        doc,
        "Compensation shall follow JIRAH's Technical Contractor Model, as "
        "detailed in Appendix B. The Consultant is engaged at an hourly rate "
        "of two hundred Canadian dollars ($200.00 CAD) per hour. All work "
        "requests are subject to a scope-and-estimate gate: the Consultant "
        "shall return a written scope summary and hour estimate before "
        "committing billable time, and shall not exceed an approved estimate "
        "by more than twenty percent (20%) without the Company's prior "
        "written approval. Invoices are submitted monthly and paid net "
        "fifteen (15) days from invoice date.",
    )

    # ---------- 4. Independent Contractor Status ----------
    add_heading(doc, "4. Independent Contractor Status")
    add_para(
        doc,
        "The Consultant acknowledges that they are engaged as an independent "
        "contractor and not as an employee, agent, or partner of the "
        "Company. The Consultant is solely responsible for payment of all "
        "taxes, insurance, benefits, GST/HST registration where applicable, "
        "and any other statutory obligations arising from this Agreement. "
        "The Consultant sets their own working hours subject to project "
        "deadlines, supplies their own tools and equipment (except for "
        "Company-issued credentials and access to client systems), and is "
        "free to engage with other clients, subject to the confidentiality, "
        "non-solicitation, and conflict-of-interest obligations of this "
        "Agreement.",
    )

    # ---------- 5. Intellectual Property ----------
    add_heading(doc, "5. Intellectual Property")
    add_para(
        doc,
        "All work product created by the Consultant in the performance of "
        "the Services — including but not limited to source code, "
        "architectural designs, infrastructure-as-code definitions, "
        "configuration, documentation, runbooks, security reviews, test "
        "suites, deployment artifacts, and any derivative works thereof — "
        'shall constitute the sole and exclusive property of JIRAH Growth '
        "Consulting Ltd. The Consultant hereby assigns to the Company, "
        "irrevocably and in perpetuity, all right, title, and interest in "
        "and to such work product upon creation, including all intellectual "
        "property rights worldwide. The Consultant waives all moral rights "
        "in such work product to the fullest extent permitted by law.",
    )
    add_para(
        doc,
        "The Consultant shall retain ownership of any pre-existing tools, "
        "libraries, frameworks, or general-purpose code created prior to or "
        'outside the scope of this Agreement ("Pre-Existing IP"), provided '
        "such Pre-Existing IP is disclosed to the Company in writing before "
        "integration into Company or client work. The Consultant grants the "
        "Company a perpetual, irrevocable, worldwide, royalty-free, "
        "transferable, sub-licensable license to use, modify, and "
        "redistribute such Pre-Existing IP solely as embedded within "
        "Company or client deliverables.",
    )
    add_para(
        doc,
        "The Consultant shall disclose in writing all third-party "
        "open-source software dependencies introduced into Company or "
        "client deliverables and shall ensure that all such dependencies "
        "are license-compatible with the Company's commercial use and "
        "downstream client deployment. The Consultant shall not introduce "
        "any undisclosed access mechanism, backdoor, telemetry, or "
        "phone-home functionality into any Company or client system.",
    )

    # ---------- 6. Confidentiality ----------
    add_heading(doc, "6. Confidentiality")
    add_para(
        doc,
        "The Consultant agrees to maintain strict confidentiality regarding "
        "all proprietary information, trade secrets, client data, business "
        "processes, internal methodologies, financial information, and "
        "system credentials obtained during the course of this Agreement, "
        "whether received from the Company or directly from Company "
        "clients. The Consultant shall not use such information for any "
        "purpose other than performance of the Services and shall not "
        "disclose such information to any third party without the Company's "
        "prior written consent. The Consultant shall notify the Company in "
        "writing within twenty-four (24) hours of any suspected breach, "
        "unauthorized access, or loss of confidential information. This "
        "obligation shall survive termination of this Agreement for a "
        "period of five (5) years.",
    )

    # ---------- 7. Non-Solicitation ----------
    add_heading(doc, "7. Non-Solicitation")
    add_para(
        doc,
        "For a period of twenty-four (24) months following termination of "
        "this Agreement, the Consultant shall not, directly or indirectly, "
        "for the Consultant's own account or on behalf of any other "
        "person or entity:",
    )
    add_bullets(
        doc,
        [
            "Solicit, induce, divert, or attempt to solicit any client of "
            "JIRAH Growth Consulting Ltd. (including any party that was a "
            "client during the term of this Agreement or in the twenty-four "
            "(24) months preceding termination) for services competitive "
            "with the Company's business; or",
            "Solicit, induce, hire, or attempt to hire any employee, "
            "consultant, or contractor of the Company.",
        ],
    )
    add_para(
        doc,
        "The parties acknowledge that this Agreement does not impose a "
        "general non-compete obligation; the Consultant may engage with "
        "other clients subject to the confidentiality, intellectual "
        "property, and non-solicitation obligations herein.",
    )

    # ---------- 8. Security and Access ----------
    add_heading(doc, "8. Security and Access")
    add_para(
        doc,
        "The Consultant shall access client systems only through "
        "Company-issued or Company-approved credentials, with multi-factor "
        "authentication enabled on all access points. The Consultant shall "
        "not store client data on personal devices or unapproved cloud "
        "services without the Company's prior written approval. The "
        "Consultant shall report any suspected security incident, "
        "credential compromise, or unauthorized access to the Company in "
        "writing within twenty-four (24) hours of discovery. Upon "
        "termination of this Agreement, the Consultant shall promptly "
        "return or securely destroy all Company and client data, "
        "credentials, and access tokens in the Consultant's possession and "
        "shall certify such return or destruction in writing to the "
        "Company.",
    )

    # ---------- 9. Code Quality and Handover Standards ----------
    add_heading(doc, "9. Code Quality and Handover Standards")
    add_para(
        doc,
        "All deliverables intended for client handover shall meet the "
        "Company's production-readiness standards, including but not "
        "limited to:",
    )
    add_bullets(
        doc,
        [
            "A documented security review covering authentication, "
            "authorization, secrets management, input validation, data "
            "encryption at rest and in transit, and audit logging;",
            "Architecture documentation, an operational runbook, and an "
            "incident-response procedure for each deployed system;",
            "Automated test coverage and a continuous integration / "
            "continuous deployment pipeline appropriate to the engagement;",
            "A defined client knowledge-transfer protocol, including any "
            "training sessions agreed in scope.",
        ],
    )
    add_para(
        doc,
        "The Consultant shall remain available on a billable basis for a "
        "thirty (30) day post-handover stabilization window for each "
        "client deployment, subject to the standard scope-and-estimate "
        "process for any work performed during that window.",
    )

    # ---------- 10. Insurance and Liability ----------
    add_heading(doc, "10. Insurance and Liability")
    add_para(
        doc,
        "The Consultant shall maintain, at the Consultant's expense, "
        "errors and omissions (professional liability) insurance with a "
        "minimum coverage limit of one million Canadian dollars "
        "($1,000,000 CAD) per claim and in aggregate annually, and shall "
        "provide evidence of such coverage to the Company upon request.",
    )
    add_para(
        doc,
        "Except for breaches of Sections 5 (Intellectual Property), 6 "
        "(Confidentiality), or 8 (Security and Access), or for acts of "
        "gross negligence or willful misconduct, the Consultant's "
        "aggregate liability under this Agreement shall not exceed the "
        "total fees paid by the Company to the Consultant in the twelve "
        "(12) months preceding the event giving rise to such liability.",
    )

    # ---------- 11. Cultural Alignment and Conduct ----------
    add_heading(doc, "11. Cultural Alignment and Conduct")
    add_para(
        doc,
        "The Consultant agrees to conduct themselves in alignment with "
        "JIRAH's mission, vision, and core values, which emphasize "
        "collaboration, adaptability, integrity, and an entrepreneurial "
        "mindset. Professional behavior, transparency, technical rigor, "
        "and commitment to client and team success are essential "
        "expectations of this role.",
    )

    # ---------- 12. Governing Law ----------
    add_heading(doc, "12. Governing Law")
    add_para(
        doc,
        "This Agreement shall be governed by and construed in accordance "
        "with the laws of the Province of British Columbia, Canada. Any "
        "disputes arising under this Agreement shall be resolved in the "
        "courts of British Columbia.",
    )

    # ---------- 13. Entire Agreement ----------
    add_heading(doc, "13. Entire Agreement")
    add_para(
        doc,
        "This Agreement, together with Appendix A and Appendix B, "
        "constitutes the entire understanding between the parties and "
        "supersedes all prior agreements or understandings, whether "
        "written or oral, relating to the subject matter herein. Any "
        "amendment to this Agreement must be in writing and signed by "
        "both parties.",
    )

    # ---------- 14. Signatures ----------
    add_heading(doc, "14. Signatures")
    add_para(
        doc,
        "IN WITNESS WHEREOF, the parties have executed this Consultancy "
        "Services Agreement as of the Effective Date.",
    )

    add_signature_block(
        doc,
        "Joshua Marshall, Principal & Founder",
        "JIRAH Growth Consulting Ltd.",
    )

    add_signature_block(
        doc,
        "Ron Bierman",
        "Senior Technical Advisor (Contractor)",
    )

    doc.add_page_break()

    # ---------- Appendix A ----------
    add_title(doc, "Appendix A")
    add_subheading(doc, "Senior Technical Advisor (Contractor) — Role Definition")
    add_para(
        doc,
        "The Senior Technical Advisor provides senior engineering "
        "leadership to JIRAH's AI product and operating-system builds. "
        "The role's central purpose is to make sure every system JIRAH "
        "delivers to a client is secure, properly architected, and "
        "production-grade — bridging the gap between rapid AI-assisted "
        "build cycles and the engineering rigor required for client "
        "handover and long-term operation.",
    )

    add_subheading(doc, "Pillar 1 — Architectural Oversight on AI-Built Systems")
    add_bullets(
        doc,
        [
            "Review system architecture for client-deployed operating "
            "systems before they reach production.",
            "Define and enforce reference architecture patterns covering "
            "authentication, authorization, data model, multi-tenancy, "
            "integration boundaries, and deployment topology.",
            "Pair-build with Jason Lotoski on complex modules where "
            "AI-assisted development requires senior engineering "
            "judgment (e.g., distributed state, role-based access "
            "control, queueing, data isolation).",
            "Specify the cloud infrastructure stack per client "
            "engagement (Azure, Supabase, GCP, or other) and own the "
            "infrastructure-as-code definition.",
        ],
    )

    add_subheading(doc, "Pillar 2 — Security and Production-Readiness Review")
    add_bullets(
        doc,
        [
            "Conduct mandatory pre-handover security reviews against an "
            "OWASP-aligned checklist covering authentication, "
            "authorization, secrets management, input validation, "
            "encryption at rest and in transit, and audit logging.",
            "Set up CI/CD pipelines, including automated tests, "
            "dependency scanning, and controlled environment promotion.",
            "Establish monitoring and alerting baselines per deployed "
            "system, covering uptime, error rates, and performance "
            "regression detection.",
            "Author an incident-response runbook for each client "
            "deployment.",
        ],
    )

    add_subheading(doc, "Pillar 3 — Client Handover and Post-Launch Stabilization")
    add_bullets(
        doc,
        [
            "Produce architecture documentation, operational runbooks, "
            "and a client onboarding kit per deployed system.",
            "Deliver knowledge-transfer sessions to client IT and "
            "operations staff where included in scope.",
            "Remain available on a billable basis for a thirty (30) day "
            "post-handover stabilization window per deployment.",
            "Operate within the JIRAH escalation path: client → Jason → "
            "Ron, with Joshua included on commercial and client-facing "
            "matters.",
        ],
    )

    add_subheading(doc, "Out of Scope")
    add_bullets(
        doc,
        [
            "Client business-process consulting (the Jirah methodology, "
            "led by Jason and Joshua).",
            "Direct client commercial conversations, pricing, or "
            "contract negotiations (Joshua and Jason only).",
            "Personnel decisions inside client organizations.",
            "Sales and pre-sales engagement work, except where senior "
            "technical review is explicitly requested by the Company.",
        ],
    )

    add_subheading(doc, "Performance Expectations")
    add_bullets(
        doc,
        [
            "Timeliness and accuracy of scope-and-estimate responses "
            "(target: within two business days of request).",
            "Estimate adherence — billable hours within twenty percent "
            "(20%) of approved estimate without prior written approval.",
            "Quality of architecture, security, and handover deliverables "
            "as measured against the standards in Section 9.",
            "Zero undisclosed security incidents or IP-disclosure "
            "violations.",
            "Collaboration, professionalism, and adherence to JIRAH's "
            "mission and values.",
        ],
    )

    doc.add_page_break()

    # ---------- Appendix B ----------
    add_title(doc, "Appendix B")
    add_subheading(doc, "Compensation and Engagement Model")

    add_subheading(doc, "Rate")
    add_para(
        doc,
        "Two hundred Canadian dollars ($200.00 CAD) per hour for all "
        "billable work performed under this Agreement. All figures are in "
        "CAD unless otherwise specified.",
    )

    add_subheading(doc, "Scope-and-Estimate-Before-Commit Workflow")
    add_para(
        doc,
        "All work performed under this Agreement is subject to the "
        "following sequence. No billable time shall be committed before "
        "the Company has approved the estimate in writing.",
    )
    add_bullets(
        doc,
        [
            "Step 1 — Request. The Company (Jason Lotoski or Joshua "
            "Marshall) describes the work in writing (email is "
            "sufficient).",
            "Step 2 — Scope and Estimate. The Consultant returns, within "
            "two (2) business days, a written scope summary, an hour "
            "estimate, and any assumptions or dependencies.",
            "Step 3 — Approval. The Company approves the estimate in "
            "writing before the Consultant begins billable work. Email "
            "approval is sufficient.",
            "Step 4 — Execution. The Consultant performs the work and "
            "flags the Company in writing as soon as it appears the "
            "estimate will be exceeded by more than twenty percent (20%). "
            "Work in excess of an approved estimate by more than twenty "
            "percent (20%) requires prior written approval and is not "
            "billable absent such approval.",
        ],
    )

    add_subheading(doc, "Invoicing and Payment")
    add_bullets(
        doc,
        [
            "The Consultant (or the Consultant's corporation) shall "
            "issue invoices monthly, itemized by engagement and date.",
            "Invoices are payable net fifteen (15) days from invoice "
            "date.",
            "Payment shall be made in Canadian dollars by electronic "
            "transfer to an account designated by the Consultant.",
            "Reimbursable expenses (travel, third-party tooling, "
            "licensing) require prior written approval and shall be "
            "itemized on the invoice with supporting receipts.",
        ],
    )

    add_subheading(doc, "Insurance Evidence")
    add_para(
        doc,
        "The Consultant shall provide evidence of the errors and "
        "omissions insurance coverage required under Section 10 within "
        "thirty (30) days of the Effective Date and upon any subsequent "
        "request by the Company.",
    )

    add_subheading(doc, "Administrative Notes")
    add_bullets(
        doc,
        [
            "Rate review annually, with any change requiring written "
            "agreement.",
            "The Company may, at its discretion, propose a retainer or "
            "fixed-fee engagement structure for specific projects, "
            "subject to mutual written agreement.",
            "No commission, equity, or revenue-share is included in this "
            "Agreement; any such arrangement would require a separate "
            "written agreement.",
        ],
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"Wrote: {out_path}")


def main() -> None:
    out_path = (
        CORPORATE
        / "Legal & Agreements"
        / "Consultancy Services Agreement - Ron Bierman - June 2026.docx"
    )
    build(out_path)


if __name__ == "__main__":
    main()
