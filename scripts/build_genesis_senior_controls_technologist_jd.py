"""
Build the v1 Senior Controls Technologist JD as a docx, matching the locked
Genesis template (Calibri, 18pt centered title, 12pt bold all-caps section
headers, List Bullet style, navy header tables with zebra striping).

Source material:
  - Tobias-content apex-tier draft (Senior-Programmer-v1.docx) — supersedes
  - Controls-Technologist.pdf (Genesis source JD) — IC craft content

Title convention: Engineer / Technologist / Programmer are synonymous at
Genesis. "Senior Controls Technologist" is the canonical title for this
seat (incumbents: Mike, Tobias). Junior tier sits below as "Junior Controls
Technologist." The Controls Technologist mid tier is built separately
(build_genesis_controls_technologist_jd.py).

Output:
  [Genesis]/07 - Deliverables/Drafts/Senior-Controls-Technologist-v1.docx
"""

from pathlib import Path
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WS / "scripts"))
from _paths import assert_no_legacy_segment, ACTIVE_CLIENTS  # type: ignore

OUT_DIR = ACTIVE_CLIENTS / "Genesis Systems" / "07 - Deliverables" / "Drafts"
OUT_PATH = OUT_DIR / "Senior-Controls-Technologist-v1.docx"
assert_no_legacy_segment(OUT_PATH)

EM = "—"
EN = "–"
LDQ = chr(0x201C)
RDQ = chr(0x201D)


def set_cell_shading(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="BFBFBF", sz="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), sz)
        b.set(qn("w:color"), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(18)


def subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(12)


def subsection(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(11)


def body(doc, text, justify=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_runs(p, text)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    _add_runs(p, text)


def _add_runs(paragraph, text):
    parts = text.split("**")
    bold_flag = False
    for chunk in parts:
        if chunk == "":
            bold_flag = not bold_flag
            continue
        r = paragraph.add_run(chunk)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.bold = bold_flag
        bold_flag = not bold_flag


def make_table(doc, headers, rows, col_widths_in=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = False
    if col_widths_in:
        for i, w in enumerate(col_widths_in):
            for cell in table.columns[i].cells:
                cell.width = Inches(w)

    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, "1F2937")
        set_cell_borders(cell, color="1F2937", sz="6")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for ri, row in enumerate(rows):
        zebra = (ri % 2 == 1)
        for ci, val in enumerate(row):
            cell = table.rows[1 + ri].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            _add_runs(p, val)
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.name = "Calibri"
            if zebra:
                set_cell_shading(cell, "F4F1EA")
            set_cell_borders(cell, color="D9D2C4", sz="4")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def build():
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = Inches(1.1)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # Footer: Page X of Y
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run("Page ")
    fr.font.size = Pt(9)
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    fr2 = footer_p.add_run()
    fr2._r.append(fld_begin); fr2._r.append(instr); fr2._r.append(fld_end)
    footer_p.add_run(" of ").font.size = Pt(9)
    fld_begin2 = OxmlElement("w:fldChar"); fld_begin2.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText"); instr2.text = "NUMPAGES"
    fld_end2 = OxmlElement("w:fldChar"); fld_end2.set(qn("w:fldCharType"), "end")
    fr3 = footer_p.add_run()
    fr3._r.append(fld_begin2); fr3._r.append(instr2); fr3._r.append(fld_end2)

    # === HEADLINE ===
    title(doc, "SENIOR CONTROLS TECHNOLOGIST")
    subtitle(doc, "Genesis Building Controls  |  Engineering & Programming Function  |  Cross-Regional Technical Track")

    # === ABOUT THE ROLE ===
    section(doc, "About the Role")
    body(doc,
         "The Senior Controls Technologist is the technical depth seat at the centre of Genesis Building "
         "Controls' delivery engine. The role owns the design, programming, integration, and commissioning of "
         "HVAC control systems on Genesis's most technically complex projects across every region. It is the "
         "seat that decides how a building behaves once it is handed over, and how cleanly Genesis's controls "
         "work travels from one project to the next.")
    body(doc,
         "This is the **apex of Genesis's controls-technologist track**. Controls Technologist and Junior "
         "Controls Technologist seats sit below it " + EM + " defined in their own role descriptions " + EM + " "
         "and the Senior Controls Technologist is the bench-builder for those tiers as Genesis hires across "
         "Kelowna, Metro Vancouver, and Alberta.")
    body(doc,
         "The role reports to the Regional Manager " + EM + " Okanagan in their capacity as Genesis's current "
         "cross-regional Technical Lead. As Genesis formalizes the Chief Engineer / Technical Lead seat "
         "alongside the Alberta launch (target horizon Q2 2027), the reporting line transitions to that seat. "
         "The evolution is named in the role's design, not held as an open question.")
    body(doc,
         "This JD formalizes work that has been carried at Genesis for years. The seat existed in practice "
         "before it had a name; the elevation gives the function a defined home, an authority profile that "
         "matches the responsibility, and a bench Genesis can develop technologists into.")

    # === ABOUT GENESIS BUILDING CONTROLS ===
    section(doc, "About Genesis Building Controls")
    body(doc,
         "Founded in 2013 inside a large electrical contracting firm and operating as an independent entity "
         "since 2017, Genesis Building Controls is one of British Columbia's fastest-growing controls and "
         "building integration companies. With expansion now active in Metro Vancouver and preparing for "
         "Alberta, Genesis serves hi-rise, commercial, institutional, residential, industrial, and healthcare "
         "clients. The firm's growth has been built on long relationships with general contractors, consulting "
         "engineers, and building owners across BC.")

    section(doc, "Core Values")
    body(doc,
         "Outstanding customer service through professional excellence, guaranteed workmanship, innovative "
         "technology, competitive pricing, and efficient scheduling " + EM + " held to the highest standards "
         "of safety and environmental protection. Integrity, dedication, and ethical practice with clients, "
         "suppliers, and staff.")

    # === REPORTING & RELATIONSHIPS ===
    section(doc, "Reporting & Relationships")
    body(doc,
         "**Reports to:** Regional Manager " + EM + " Okanagan (current cross-regional Technical Lead). "
         "Transitions to Chief Engineer / Technical Lead upon formalization of that seat (target horizon "
         "Q2 2027 alongside the Alberta launch).",
         justify=False)
    body(doc,
         "**Direct reports:** Controls Technologists and Junior Controls Technologists across all Genesis "
         "regions. Hiring, onboarding, performance conversations, and bench development for the controls-"
         "technologist pool sit with this role in partnership with the Technical Lead and the Office Manager.",
         justify=False)
    body(doc,
         "**Cross-regional functional authority:** The Senior Controls Technologist sets controls design "
         "standards, programming conventions, sequence-of-operations methodology, code-review discipline, and "
         "platform selection across all Genesis regions. Technologists in Kelowna, Metro Vancouver, and "
         "Alberta work to the same standard " + EM + " set, maintained, and continuously raised at this seat.",
         justify=False)

    subsection(doc, "Primary internal partners")
    bullet(doc, "**Regional Manager " + EM + " Okanagan** " + EM + " direct reporting line; the Senior Controls Technologist is the day-to-day execution of the cross-regional Technical Lead function the RM-Okanagan currently carries.")
    bullet(doc, "**Regional Manager " + EM + " Metro Vancouver** " + EM + " partners on Vancouver-region project controls priorities, technician scheduling around commissioning windows, and field escalation routing.")
    bullet(doc, "**Lead Estimator** " + EM + " partners at bid-stage on technically demanding pursuits " + EM + " sequence-of-operations complexity, integration scope, and controls effort that drives bid pricing.")
    bullet(doc, "**Service & Warranty Coordinator** " + EM + " partners on service-side programming, software maintenance agreement renewals, retrofit work, and post-handover code maintenance.")
    bullet(doc, "**Office Manager** " + EM + " partners on controls-team hiring, software license administration, and contract administration for integration partners and tool vendors.")
    bullet(doc, "**Project Managers (all regions)** " + EM + " day-to-day project delivery; the Senior Controls Technologist supports complex commissioning windows and resolves technical escalations the PMs surface.")

    # === CADENCE & SUPPORT ===
    section(doc, "Cadence & Support Commitments")
    body(doc,
         "The Senior Controls Technologist is the technical anchor of the engineering and programming "
         "function. The commitments to the incumbent are part of the role's design, not optional add-ons:")
    bullet(doc, "**Weekly 1:1 with the Technical Lead** (currently the Regional Manager " + EM + " Okanagan) " + EM + " 45 minutes, recurring, agenda-driven. Active project controls work, blockers, standards, bench development.")
    bullet(doc, "**Weekly engineering sync** " + EM + " Technical Lead + Senior Controls Technologist + Controls Technologists " + EM + " structured cadence covering active project programming, sequence-of-operations review, and integration milestones.")
    bullet(doc, "**Monthly engineering scorecard review** " + EM + " Senior Controls Technologist presents controls-function performance (project on-time / on-spec, commissioning defect rate, bench development) to the Technical Lead and Owner.")
    bullet(doc, "**Quarterly leadership review** " + EM + " Senior Controls Technologist attends as the engineering & programming voice in the senior operational leadership review (both Regional Managers, Lead Estimator, Office Manager, Senior Controls Technologist, Owner) at Kelowna headquarters.")
    bullet(doc, "**Standing first question:** " + LDQ + "How can I support you in this role? What do you need from me?" + RDQ + " The Technical Lead asks it of the Senior Controls Technologist; the Senior Controls Technologist asks it of each direct report. This is a Genesis cultural primitive.")

    # === KEY RESPONSIBILITIES ===
    section(doc, "Key Responsibilities")

    subsection(doc, "HVAC Controls Design & Engineering")
    bullet(doc, "Design and develop HVAC control systems on Genesis's most technically complex projects based on project specifications, building codes, and industry standards.")
    bullet(doc, "Perform detailed site surveys and assessments to gather the information necessary for system design and implementation.")
    bullet(doc, "Select appropriate control components, sensors, actuators, and other hardware necessary for system integration.")
    bullet(doc, "Develop control strategies and sequences of operation to optimize HVAC system performance and energy efficiency.")
    bullet(doc, "Create control system drawings, schematics, and diagrams using CAD software.")
    bullet(doc, "Conduct (or oversee) sequence-of-operations review on every flagship project before programming starts " + EM + " catch ambiguities, drawing errors, or consulting-engineer assumptions that would cost field time to discover.")

    subsection(doc, "Programming & Integration")
    bullet(doc, "Program and configure building automation systems (BAS) and control devices " + EM + " PLCs, DDC controllers, thermostats " + EM + " on Genesis's most complex projects.")
    bullet(doc, "Own the integration logic between Genesis's controls work and external systems " + EM + " third-party hardware, building management overlays, mechanical and life-safety panels, owner-side dashboards.")
    bullet(doc, "Stand up the programming on new platforms, controller families, and protocol families as Genesis takes on work that requires them; document the pattern so other technologists can run it next time.")

    subsection(doc, "Commissioning & Field Liaison")
    bullet(doc, "Support commissioning windows " + EM + " including on-site presence on tight or high-stakes commissioning days " + EM + " and own programming-side resolution of commissioning issues.")
    bullet(doc, "Conduct system testing, commissioning, and troubleshooting to ensure proper functionality and performance.")
    bullet(doc, "Be the senior controls voice on field escalations " + EM + " both as the routing point and as the answer when the question is at the depth this seat carries.")
    bullet(doc, "**Triage discipline:** routine controls questions route to the Controls Technologist pool; controller-specific issues to the in-region technologist; commissioning-day support to whichever technologist owns the project. The Senior Controls Technologist engages on hard escalations, novel integrations, and cross-regional pattern questions " + EM + " not routine work that belongs elsewhere.")
    bullet(doc, "Lead post-commissioning defect review on each project " + EM + " feed findings back into the controls standards and the design library.")
    bullet(doc, "Provide technical support and assistance to field technicians and subcontractors during installation, startup, and maintenance activities.")

    subsection(doc, "Cross-Regional Standards & Code Quality")
    bullet(doc, "Set and maintain Genesis's controls design and programming standards " + EM + " sequence-of-operations templates, controller programming conventions, BACnet object naming, graphics standards, alarming and trending discipline. Standards travel across regions; an Alberta job runs on the same logic as a Kelowna job.")
    bullet(doc, "Run code review on programming work issued by Controls Technologist and Junior Controls Technologist seats " + EM + " both as a quality gate and as a coaching mechanism. Code review is the developmental spine of the bench.")
    bullet(doc, "Maintain the controls library " + EM + " reusable logic blocks, design templates, integration patterns, commissioning checklists " + EM + " so that knowledge compounds across projects instead of being re-derived each time.")
    bullet(doc, "Lead platform selection and major version upgrade decisions (Niagara N4, controller families, integration middleware). Joint with the Technical Lead and Owner where the choice carries commercial or capital impact.")

    subsection(doc, "Bid-Stage Technical Input")
    bullet(doc, "Partner with the Lead Estimator at bid-stage on technically demanding pursuits " + EM + " confirm sequence-of-operations scope, identify controls effort that drives pricing, and name integration unknowns that need to be flagged in the bid.")
    bullet(doc, "Attend bid-strategy review on pursuits where controls complexity is the principal driver of project margin.")
    bullet(doc, "Provide go / no-go input on pursuits whose controls scope is outside Genesis's current capability, or whose integration risk is misaligned with the firm's bench capacity.")

    subsection(doc, "Bench Leadership")
    bullet(doc, "Lead, coach, and develop the controls-technologist bench " + EM + " Controls Technologists and Junior Controls Technologists " + EM + " across all regions. Set the bar for what controls quality looks like at Genesis and raise the bench to that bar.")
    bullet(doc, "Own hiring, onboarding, and performance conversations for the controls-technologist function in partnership with the Technical Lead and Office Manager.")
    bullet(doc, "Run weekly engineering sync as the technical chair " + EM + " structured agenda, blocker review, standards updates, knowledge transfer.")
    bullet(doc, "Conduct quarterly 1:1 development conversations with each direct report; formal annual performance review by April 30 each year.")
    bullet(doc, "Build the bench Genesis needs to scale " + EM + " succession planning for the function as the firm grows toward $10M revenue and the Alberta launch lands.")

    # === AUTHORITY MATRIX ===
    section(doc, "Authority Matrix")
    body(doc,
         "Numerical thresholds are confirmed in the offer letter and calibrated against the engineering & "
         "programming function's annual budget.")

    auth_rows = [
        ["Controls standards, code review, sequence-of-operations methodology",
         "Full " + EM + " sets the firm standard",
         "None"],
        ["Project controls design and programming approach, controller selection within approved platforms",
         "Full",
         "Outside approved platforms " + EN + " Technical Lead"],
        ["Controls library and template changes",
         "Full",
         "None"],
        ["Platform / major version upgrade decisions (Niagara, controller families)",
         "Joint with Technical Lead",
         "Capital impact > **$" + "{X}" + "** " + EN + " Owner"],
        ["Software license procurement within approved budget",
         "Up to **$" + "{X}" + "** per single license / per annual contract",
         "> threshold " + EN + " Technical Lead + Office Manager"],
        ["Hiring within approved controls headcount (Controls Technologist, Junior Controls Technologist)",
         "Full",
         "New roles outside plan " + EN + " Technical Lead + Owner"],
        ["Termination & disciplinary decisions, controls staff",
         "Joint with Technical Lead and Owner",
         "All cases consult Technical Lead first"],
        ["Bid-stage go / no-go input on controls-driven pursuits",
         "Advisory " + EM + " recommendation to Lead Estimator + Owner",
         "n/a (final call sits with Lead Estimator + Owner)"],
        ["Field escalation resolution within approved scope",
         "Full",
         "Cross-regional capacity conflict " + EN + " Technical Lead + RM"],
        ["External integration partner / contractor engagement",
         "Up to **$" + "{X}" + "** per engagement, within approved budget",
         "> threshold " + EN + " Technical Lead + Owner"],
        ["Capital tooling investment beyond approved budget",
         "Owner-owned",
         "All " + EM + " Senior Controls Technologist recommends"],
        ["Material legal, financial, or reputational exposure",
         "Owner-owned",
         "Immediate notification"],
    ]
    make_table(
        doc,
        ["Decision Area", "Senior Controls Technologist Authority", "Escalation Trigger"],
        auth_rows,
        col_widths_in=[2.8, 2.3, 1.6],
    )

    # === SUCCESS METRICS ===
    section(doc, "Success Metrics")
    body(doc, "The Senior Controls Technologist's performance is reviewed monthly against:")
    bullet(doc, "**Project controls on-time / on-spec rate** " + EM + " controls-side delivery against project schedule and commissioning targets.")
    bullet(doc, "**Commissioning defect rate** " + EM + " rework attributable to controls work, measured by project, by region, and by controller family. Trend over rolling quarter.")
    bullet(doc, "**Sequence-of-operations review coverage** " + EM + " share of flagship projects where SOO review happened before programming started; misses tracked.")
    bullet(doc, "**Standards adoption** " + EM + " share of controls work shipped through the standards / template / library system, versus one-off.")
    bullet(doc, "**Bid-stage technical-input turnaround** " + EM + " time from Lead Estimator request to controls-effort estimate returned.")
    bullet(doc, "**Bench retention and progression** " + EM + " direct-report retention, demonstrable progression on the Controls Technologist / Junior Controls Technologist ladder, and successful onboarding cycle time for new controls hires.")
    bullet(doc, "**Field escalation volume to this seat** " + EM + " measured monthly. Routine escalations are a leading indicator of standards or bench-depth gaps; persistent high volume gets named.")

    # === TRANSITION ARC ===
    section(doc, "Transition Arc " + EM + " First 90 Days")
    body(doc,
         "This role is being filled by an internal formalization. The incumbent has carried this work for "
         "years; the JD names the seat and gives it a defined home. The first 90 days are about making the "
         "seat visible " + EM + " installing the cadence, naming the standards, and starting the bench-"
         "leadership rhythm.")

    subsection(doc, "First 30 Days")
    bullet(doc, "Weekly 1:1 with the Technical Lead installed and running.")
    bullet(doc, "Weekly engineering sync agenda formalized " + EM + " recurring, structured, on-calendar.")
    bullet(doc, "Controls standards inventory " + EM + " what is written down today, what lives in heads, what needs to be documented next.")
    bullet(doc, "Initial bench check-in completed with each direct report using the " + LDQ + "How can I support you?" + RDQ + " frame.")

    subsection(doc, "Days 30" + EN + "60")
    bullet(doc, "First wave of controls standards documented and shared across the bench " + EM + " sequence-of-operations templates, BACnet object naming conventions, alarming and trending discipline.")
    bullet(doc, "Controls library structure stood up " + EM + " reusable logic blocks and design templates indexed and accessible to the bench.")
    bullet(doc, "Code-review cadence installed for Controls Technologist and Junior Controls Technologist work.")
    bullet(doc, "First monthly engineering scorecard delivered to the Technical Lead and Owner.")

    subsection(doc, "Days 60" + EN + "90")
    bullet(doc, "First quarterly engineering-function review delivered to the senior operational leadership team.")
    bullet(doc, "Controls standards adoption rate measured and reported across regions.")
    bullet(doc, "Controls Technologist / Junior Controls Technologist development plans drafted for each direct report.")
    bullet(doc, "At least one cross-regional controls decision held by the Senior Controls Technologist in the first 90 days that the Technical Lead or Owner would have stepped into a year ago.")

    # === COMPENSATION FRAMEWORK ===
    section(doc, "Compensation Framework")
    bullet(doc, "**Structure:** base salary plus performance variable. Structure consistent with senior technical roles at Genesis; numerical alignment calibrated by function and tenure.")
    bullet(doc, "**Variable tied to:** controls-side delivery performance (on-time / on-spec rate), commissioning defect rate, and bench retention and development. Specific weighting confirmed at structure finalization.")
    bullet(doc, "**Vehicle / phone / professional development:** included; specifics confirmed at structure finalization. Professional development allowance covers certifications (Tridium / Niagara, BACnet, integration platforms) and conference attendance.")
    bullet(doc, "**Review cadence:** annual, April; informal mid-year touch in October.")
    bullet(doc, "**Long-term incentive eligibility:** as Genesis formalizes a multi-year incentive plan, this role is eligible.")
    bullet(doc, "**Final compensation structure locked:** within 14 days of role acceptance.")

    # === PERFORMANCE REVIEW ===
    section(doc, "Performance Review & Role Confirmation")
    bullet(doc, "**30-day check-in** " + EM + " Technical Lead and Senior Controls Technologist review transition deliverables (cadence installed, standards inventory complete, bench check-ins done).")
    bullet(doc, "**90-day check-in** " + EM + " Technical Lead and Senior Controls Technologist review the full transition arc against milestones; first quarterly engineering review reviewed.")
    bullet(doc, "**Six-month role confirmation review** " + EM + " formal review at the end of December. Outcome is one of: role confirmed (default expectation), role confirmed with development plan, role not yet confirmed " + EM + " extended review. This is a fairness mechanism, not a probationary trap. Both parties walk in with prepared input.")
    bullet(doc, "**Annual performance review** " + EM + " every April, against the full Success Metrics scorecard.")

    # === QUALIFICATIONS ===
    section(doc, "Qualifications & Experience")

    subsection(doc, "Required")
    bullet(doc, "Minimum 8 years of HVAC building controls design and programming experience, with demonstrated depth across multiple controller families and the dominant North American controls platforms.")
    bullet(doc, "Bachelor's degree in Mechanical Engineering, Electrical Engineering, or related field " + EM + " or equivalent technologist credential (CET / RET) backed by depth of practice.")
    bullet(doc, "Deep working knowledge of Tridium Niagara (AX and N4), BACnet (MS/TP and IP), Modbus, LonWorks, and at least one major manufacturer controller stack (Distech, Honeywell, Johnson Controls, Schneider Electric, or equivalent).")
    bullet(doc, "Strong knowledge of HVAC systems " + EM + " air handling units, chillers, boilers, VAV systems " + EM + " and integration with mechanical, electrical, and life-safety systems as it touches controls.")
    bullet(doc, "Proficiency in CAD software for control system design and documentation.")
    bullet(doc, "Demonstrated track record of sequence-of-operations design " + EM + " from consulting-engineer drawings through to commissioned, owner-accepted programming.")
    bullet(doc, "Demonstrated ability to lead and coach other technologists " + EM + " set standards, run code review, develop the bench.")
    bullet(doc, "Proficient in at least one programming language commonly used in HVAC controls (ladder logic, C/C++, Java, Python, or equivalent).")
    bullet(doc, "Strong written and verbal communication; ability to translate between technologist, engineer, field, and owner-side audiences.")
    bullet(doc, "Excellent problem-solving skills and rigorous attention to detail on logic, alarming, trending, and graphics.")
    bullet(doc, "Valid BC driver's license and access to reliable transportation.")

    subsection(doc, "Preferred")
    bullet(doc, "Master's degree in Mechanical or Electrical Engineering.")
    bullet(doc, "Niagara certifications (TCP, certified instructor, or equivalent senior-level credentials).")
    bullet(doc, "Professional certifications such as Certified Energy Manager (CEM) or Certified Automation Professional (CAP).")
    bullet(doc, "Experience leading a controls bench across multiple geographies or business units.")
    bullet(doc, "Familiarity with the BC and Alberta consulting engineer and general contractor markets.")
    bullet(doc, "Experience with controls cybersecurity and network-segmentation discipline in commercial buildings.")

    # === BEHAVIOURAL COMPETENCIES ===
    section(doc, "Behavioural Competencies")
    body(doc, "These are the behaviours expected in role and the spine of every coaching 1:1.")
    comp_rows = [
        ["**Writes for the next person who reads the code**",
         "Logic is readable. Names are honest. Graphics, alarming, and trending follow the standard. The next technologist to touch the project " + EM + " six months or six years from now " + EM + " can find their footing without an interpreter.",
         "Clever shortcuts that compile and ship. Heroic individual work that no one else can pick up. " + LDQ + "I'll explain it later" + RDQ + " code that becomes load-bearing."],
        ["**Sets the standard, not just the deliverable**",
         "Controls standards are documented, current, and used. New technologists onboard against the documented standard, not against tribal knowledge. Standards travel across regions.",
         "Standards live in one head. Each technologist improvises. Reusable logic gets re-derived per project because no one has owned the library."],
        ["**Coaches the bench rather than catching the work**",
         "Walks a junior technologist through how to think about a sequence or an integration instead of writing it for them. Code review is teaching, not gatekeeping. Develops the bench by raising it, not by carrying it.",
         LDQ + "It's faster if I do it." + RDQ + " Defaults to writing the code personally because the depth makes it easy. Absorbs work that should route through the Controls Technologist / Junior Controls Technologist seats."],
        ["**Calm under commissioning pressure**",
         "Commissioning day lands; the team sees the same demeanor as on a normal Tuesday. Filters noise before it reaches the field crew. Solves the actual problem instead of the visible one.",
         "Visible reactivity under pressure. Pulls the bench into firefighting by tone. Lets phone-call frequency dictate the day."],
        ["**Holds the line on sequence-of-operations integrity**",
         "Names ambiguous drawings, missing assumptions, or consulting-engineer gaps before programming starts " + EM + " even when the project clock is tight. Surfaces a flag in the next 1:1 rather than burying it in commissioning rework.",
         "Programs to the drawings as drawn and lets the gaps land at commissioning. " + LDQ + "We'll figure it out on site" + RDQ + " attitude. Treats sequence-of-ops review as paperwork rather than risk-reduction."],
        ["**Reads the field, not just the spec**",
         "Knows what the controls have to do in a real building under real load " + EM + " not just what the drawings say they should do. Walks the job at least once before programming, more for the complex ones.",
         "Treats the spec as gospel. Programs in isolation from the mechanical reality. Surprises at commissioning that an hour on site would have prevented."],
        ["**Owner-mindset on platform choice**",
         "Brings recommendations on platform, version, and tooling investment with a clear " + LDQ + "here's what it costs, here's what it buys, here's what it commits us to for the next five years" + RDQ + " frame. Reads the firm's cycle and adjusts.",
         "Treats platform choice as a technologist's-only preference. Doesn't translate technical choice into commercial implication. Walks into Owner conversations with open questions instead of recommendations."],
    ]
    make_table(
        doc,
        ["Competency", "Looks Like", "Does Not Look Like"],
        comp_rows,
        col_widths_in=[1.7, 2.5, 2.5],
    )

    # === WORK ENVIRONMENT & TRAVEL ===
    section(doc, "Work Environment & Travel")
    bullet(doc, "Based at the assigned regional office (Kelowna headquarters or Metro Vancouver, by incumbent).")
    bullet(doc, "Regular jobsite presence across BC for sequence-of-operations review, commissioning support, and complex-project programming days. Cadence calibrated to project mix.")
    bullet(doc, "Monthly travel to the counterpart regional office for cross-regional engineering sync, commissioning days, and consulting-engineer engagement " + EM + " typically 1" + EN + "2 days per visit.")
    bullet(doc, "Quarterly extended leadership offsite (Genesis senior operational leadership team) at Kelowna.")
    bullet(doc, "Travel to Alberta as the regional launch establishes a controls presence; cadence calibrated to project pace and on-site technologist onboarding.")
    bullet(doc, "Industry event attendance " + EM + " Tridium / Niagara user groups, BACnet conferences, and equivalent technical forums.")

    # === GROWTH PATHWAY ===
    section(doc, "Growth Pathway")
    body(doc,
         "The Senior Controls Technologist role sits on Genesis's technical leadership track. The clear next "
         "step is the **Chief Engineer / Technical Lead** seat as it formalizes alongside the Alberta launch "
         "(target horizon Q2 2027). Strong performance in this role " + EM + " particularly the bench-building, "
         "the cross-regional standards work, and the technical leadership of Genesis's most complex projects "
         + EM + " positions the incumbent as a leading candidate for that seat. The role's depth across "
         "controls design, programming, and integration makes it one of the few seats at Genesis that touches "
         "every dimension of the firm's technical posture as it scales past $10M revenue.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    p = build()
    print(f"Wrote: {p}")
